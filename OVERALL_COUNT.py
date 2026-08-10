# =========================
# IMPORTS
# =========================
import pandas as pd
import pyreadstat as py
import re
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pathlib import Path 



# =========================
# LOAD DATA
# =========================
#
File_path=Path.home()/"Desktop"/"Python check"/"DATASET"/"Spss"
retail_path=File_path/"PRE FINAL RISE.sav"

df, meta = py.read_sav(retail_path)
df.columns = df.columns.str.strip()


# MULTI MAPS (MULTIPLE)
# =========================
multi_maps = {
   
}


# Custom headers dictionary
custom_headers = {
    # --- Professional Background & Experience ---
}

# =========================
# APPLY LABELS
# =========================
def apply_labels(df, meta, var):

    label_name = meta.variable_to_label.get(var)
    labels_dict = meta.value_labels.get(label_name)

    col = f"{var}_label"

    if labels_dict is not None:
        df[var] = df[var].astype("Int64")
        df[col] = df[var].map(labels_dict)
    else:
        df[col] = df[var].astype(str)

    return df, labels_dict

# to return label instead of variable

def get_var_label(meta, var):
    return meta.column_names_to_labels.get(var, var)

# =========================
# SINGLE TABLE
# =========================
def single_table(df, meta, row, var):

    df, labels_dict = apply_labels(df, meta, var)

    table = pd.crosstab(
        df[f"{row}_label"],
        df[f"{var}_label"],
        dropna=False
    )

    if labels_dict is not None:
        table = table.reindex(columns=list(labels_dict.values()), fill_value=0)

    var_label = get_var_label(meta, var)
    table.columns = pd.MultiIndex.from_product([[var_label], table.columns])
    row_label = get_var_label(meta, row)
    table.index.name = row_label

    return table


# =========================
# MULTI TABLE (COLUMN MODE)
# =========================
def multi_table(df, row, var, multi_map):

    split_df = df[var].astype(str).str.get_dummies(sep=" ")
    split_df = split_df.rename(columns=multi_map)
    split_df = split_df.reindex(columns=list(multi_map.values()), fill_value=0)

    table = split_df.groupby(df[f"{row}_label"]).sum()

    var_label = get_var_label(meta, var)
    table.columns = pd.MultiIndex.from_product([[var_label], table.columns])
    row_label = get_var_label(meta, row)
    table.index.name = row_label

    return table


# =========================
# MASTER FUNCTION (NORMAL)
# =========================
def create_table(df, meta, row, single_cols, multi_cols, multi_maps, stats_var=None):

    df, row_labels = apply_labels(df, meta, row)

    # FIX: capture row question label early so we can reapply it after concat
    row_label = get_var_label(meta, row)

    tables = []

    # =========================
    # SINGLE VARIABLES
    # =========================
    for var in single_cols:
        tables.append(single_table(df, meta, row, var))

    # =========================
    # MULTI VARIABLES
    # =========================
    for var in multi_cols:
        var_map = multi_maps.get(var, {})
        tables.append(multi_table(df, row, var, var_map))

    final_table = pd.concat(tables, axis=1)

    if row_labels is not None:
        final_table = final_table.reindex(list(row_labels.values()))

    final_table = final_table.fillna(0).replace(0, "-")

    # =========================
    # BASE ROW
    # =========================
    base_row = final_table.replace("-", 0).astype(float).sum(axis=0)
    base_row = pd.DataFrame([base_row], index=["Base"])
    # FIX: carry the question label onto base_row so concat preserves index.name
    base_row.index.name = row_label

    final_table = pd.concat([base_row, final_table])

    # =========================
    # ADD MEAN / SD / SE (FIXED)
    # =========================
    if stats_var:

        df[stats_var] = pd.to_numeric(df[stats_var], errors="coerce")

        stats_blocks = []

        var_label_map = {get_var_label(meta, v): v for v in single_cols}

        for group in final_table.columns.get_level_values(0).unique():

            if group == "":
                continue

            sub_cols = final_table[group].columns

            temp_stats = pd.DataFrame(
                index=["Mean", "Standard deviation", "Standard error"],
                columns=sub_cols
            )

            # -------------------------
            # SINGLE VARIABLE
            # -------------------------
            if group in var_label_map:

                real_var = var_label_map[group]

                df, labels_dict = apply_labels(df, meta, real_var)

                grouped = df.groupby(f"{real_var}_label")[stats_var]

                mean = grouped.mean()
                std = grouped.std()
                count = grouped.count()
                se = std / (count ** 0.5)

                temp_stats.loc["Mean"] = mean.reindex(sub_cols)
                temp_stats.loc["Standard deviation"] = std.reindex(sub_cols)
                temp_stats.loc["Standard error"] = se.reindex(sub_cols)

            # -------------------------
            # MULTI VARIABLE
            # -------------------------
            else:

                for var, mp in multi_maps.items():

                    if get_var_label(meta, var) == group:

                        split_df = df[var].astype(str).str.get_dummies(sep=" ")
                        split_df = split_df.rename(columns=mp)

                        for col in sub_cols:

                            if col in split_df.columns:
                                sub = df.loc[split_df[col] == 1, stats_var]
                            else:
                                sub = pd.Series()

                            sub = sub.dropna()

                            if len(sub) > 0:
                                mean = sub.mean()
                                std = sub.std()
                                se = std / (len(sub) ** 0.5)
                            else:
                                mean, std, se = 0, 0, 0

                            temp_stats.loc["Mean", col] = round(mean, 2)
                            temp_stats.loc["Standard deviation", col] = round(std, 2)
                            temp_stats.loc["Standard error", col] = round(se, 2)

            temp_stats.columns = pd.MultiIndex.from_product([[group], temp_stats.columns])
            stats_blocks.append(temp_stats)

        stats_full = pd.concat(stats_blocks, axis=1)

        stats_full = stats_full.reindex(columns=final_table.columns, fill_value="")

        final_table = pd.concat([final_table, stats_full])

    # =========================
    #  FINAL TOTAL (CORRECT FOR BOTH COUNT + STATS)
    # =========================
    numeric_table = final_table.replace("-", 0)

    stats_rows = ["Mean", "Standard deviation", "Standard error"]

    total_values = []

    # clean stats variable once
    if stats_var:
        stats_data = df[stats_var].dropna()

    for idx in final_table.index:

        if stats_var and idx == "Mean":
            total_values.append(round(stats_data.mean(), 2))

        elif stats_var and idx == "Standard deviation":
            total_values.append(round(stats_data.std(), 2))

        elif stats_var and idx == "Standard error":
            total_values.append(round(stats_data.std() / (len(stats_data) ** 0.5), 2))

        else:
            total_values.append(
                numeric_table.loc[idx].astype(float).sum()
            )

    total_df = pd.DataFrame(total_values, index=final_table.index)
    total_df.columns = pd.MultiIndex.from_tuples([("", "Total")])

    final_table = pd.concat([total_df, final_table], axis=1)

    # FIX: reapply question label as index name after all concats
    final_table.index.name = row_label

    return final_table

# =========================
# MULTI AS ROW (FULLY UPGRADED)
# =========================

def multi_row_table(df, meta, row_var, col_vars, multi_maps, header_dict=None):
    """
    Supports:
    ✔ multi as row
    ✔ single as column
    ✔ multi as column
    ✔ Custom row headers from dictionary
    """

    # FIX: resolve the row question label from header_dict first, then meta
    row_label = (header_dict.get(row_var) if header_dict else None) or get_var_label(meta, row_var)

    # -------------------------
    # ROW MULTI
    # -------------------------
    row_map = multi_maps.get(row_var, {})

    split_row = df[row_var].astype(str).str.get_dummies(sep=" ")
    split_row = split_row.rename(columns=row_map)
    split_row = split_row.reindex(columns=list(row_map.values()), fill_value=0)

    tables = []

    # -------------------------
    # LOOP THROUGH COLUMNS
    # -------------------------
    for col in col_vars:
        if col not in multi_maps:
            df, labels_dict = apply_labels(df, meta, col)
            temp = split_row.groupby(df[f"{col}_label"]).sum().T
            if labels_dict is not None:
                temp = temp.reindex(columns=list(labels_dict.values()), fill_value=0)
        else:
            col_map = multi_maps[col]
            split_col = df[col].astype(str).str.get_dummies(sep=" ")
            split_col = split_col.rename(columns=col_map)
            split_col = split_col.reindex(columns=list(col_map.values()), fill_value=0)
            temp = split_row.T.dot(split_col)

        col_label = get_var_label(meta, col) or col
        temp.columns = pd.MultiIndex.from_product([[col_label], temp.columns])
        tables.append(temp)

    # -------------------------
    # MERGE ALL
    # -------------------------
    final_table = pd.concat(tables, axis=1)
    final_table = final_table.fillna(0).replace(0, "-")

    # =========================
    # BASE ROW
    # =========================
    base_row_vals = final_table.replace("-", 0).astype(float).sum(axis=0)
    base_row = pd.DataFrame([base_row_vals], index=["Base"])
    # FIX: carry the question label onto base_row so concat preserves index.name
    base_row.index.name = row_label

    final_table = pd.concat([base_row, final_table])

    # =========================
    # TOTAL COLUMN (LEFT)
    # =========================
    first_group = final_table.columns.get_level_values(0)[0]
    group_df = final_table[first_group]
    total_col = group_df.replace("-", 0).astype(float).sum(axis=1)

    total_df = pd.DataFrame(total_col)
    total_df.columns = pd.MultiIndex.from_tuples([("", "Total")])

    final_table = pd.concat([total_df, final_table], axis=1)

    # FIX: reapply question label as index name after all concats
    final_table.index.name = row_label

    return final_table

def auto_multi_table(df, meta, prefix, col_vars, header_dict):
    # 1. Find all columns starting with the prefix (e.g., "Q30")
    relevant_cols = [c for c in df.columns if c.startswith(prefix)]

    # 2. Automatically build the map from the column headers
    # It takes everything after the colon ':'
    extracted_map = {
        str(i+1): c.split(':')[-1].strip()
        for i, c in enumerate(relevant_cols)
    }

    # 3. Merge the columns into the numeric string format ("1 3 5")
    merged_name = f"{prefix}_Merged"
    df[merged_name] = df[relevant_cols].apply(
        lambda x: " ".join([str(i+1) for i, v in enumerate(x) if str(v).lower() in ['yes', '1', '1.0']]),
        axis=1
    )

    # 4. Create a temporary multi_map just for this function call
    temp_map = {merged_name: extracted_map}

    # 5. Return the created table
    return multi_row_table(
        df, meta,
        row_var=merged_name,
        col_vars=col_vars,
        multi_maps=temp_map,
        header_dict=header_dict
    )



def spss_multi_table(df, meta, prefix, col_vars, header_dict, specific_cols=None):
    # 1. Identify sub-columns
    if specific_cols:
        relevant_cols = specific_cols
    else:
        # Regex: find columns that start with prefix, followed by _ and then ONLY digits
        # This excludes things like 'Q30_rank' or 'Q30_text'
        regex_pattern = re.compile(f"^{prefix}_\\d+$")
        relevant_cols = [c for c in df.columns if regex_pattern.match(c)]

    if not relevant_cols:
        print(f"Warning: No valid columns found for prefix {prefix}")
        return None

    extracted_map = {}
    main_question = ""

    # 2. Extract labels and main question
    for i, c in enumerate(relevant_cols):
        label = meta.column_names_to_labels.get(c, "")

        if label:
            # Split logic to separate Question from Option
            if ":" in label:
                parts = label.split(':', 1)
                if not main_question: main_question = parts[0].strip()
                option_text = parts[1].strip()
            elif "?" in label:
                parts = label.split('?', 1)
                if not main_question: main_question = parts[0].strip() + "?"
                option_text = parts[1].strip()
            else:
                option_text = label

            extracted_map[str(i+1)] = option_text

    # 3. Create the Merged Data Column
    merged_name = f"{prefix}_Merged"
    # Ensure we only count '1' or 'Yes'
    df[merged_name] = df[relevant_cols].apply(
        lambda x: " ".join([str(i+1) for i, v in enumerate(x) if str(v).strip().lower() in ['yes', '1', '1.0']]),
        axis=1
    )

    # 4. Update the header_dict (using the Question text found in metadata)
    header_dict[merged_name] = main_question if main_question else prefix

    # 5. Generate Table
    return multi_row_table(
        df, meta,
        row_var=merged_name,
        col_vars=col_vars,
        multi_maps={merged_name: extracted_map},
        header_dict=header_dict
    )


def format_excel_table(writer, sheet_name, current_row, df):
    workbook = writer.book
    worksheet = workbook[sheet_name]
    
    # --- Define Styles ---
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    base_row_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    border_side = Side(style='thin', color="000000")
    border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # --- Setup Dimensions ---
    header_depth = df.columns.nlevels
    index_width = df.index.nlevels
    total_cols = len(df.columns) + index_width
    data_start_row = current_row + 1 + header_depth

    # 1. Format Headers
    for r in range(current_row + 1, data_start_row):
        for c in range(1, total_cols + 1):
            cell = worksheet.cell(row=r, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border

    # 2. Format Base Row
    for c in range(1, total_cols + 1):
        cell = worksheet.cell(row=data_start_row, column=c)
        cell.fill = base_row_fill
        cell.font = Font(bold=True)
        cell.alignment = center_align if c > index_width else left_align
        cell.border = border

    # 3. Format Data Rows
    for r in range(data_start_row + 1, data_start_row + len(df)):
        for c in range(1, total_cols + 1):
            cell = worksheet.cell(row=r, column=c)
            cell.border = border
            cell.alignment = center_align if c > index_width else left_align
            
def write_table(writer, df, sheet, current_row, title=None, table_num=None):
    workbook = writer.book

    if sheet not in workbook.sheetnames:
        workbook.create_sheet(sheet)

    worksheet = workbook[sheet]

    if title is not None:
        label = f"Table {table_num}: {title}" if table_num is not None else title
        worksheet.cell(row=current_row + 1, column=1).value = label
        current_row += 1

    df.to_excel(writer, sheet_name=sheet, startrow=current_row)

    format_excel_table(writer, sheet, current_row, df)   # ← removed has_title=False

    return current_row + df.columns.nlevels + len(df) + 3


table64 = spss_multi_table(
    df, 
    meta, 
    "Q24", ["outlet_type"],
    custom_headers
    )
print(table64)